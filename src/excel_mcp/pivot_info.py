"""Utilities for reading pivot table metadata from Excel workbooks."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Iterable

import msgspec
from msgspec import json as msgjson
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .exceptions import PivotError, SheetError, ValidationError

logger = logging.getLogger(__name__)

XLSX_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}


class PivotValueInfo(msgspec.Struct, kw_only=True):
    """Metadata about a pivot value field."""

    display_name: str | None = None
    source_field: str | None = None
    aggregation: str | None = None


class PivotTableInfo(msgspec.Struct, kw_only=True):
    """Metadata about a single pivot table."""

    name: str | None = None
    location: str | None = None
    rows: list[str] = msgspec.field(default_factory=list)
    columns: list[str] = msgspec.field(default_factory=list)
    values: list[PivotValueInfo] = msgspec.field(default_factory=list)
    filters: list[str] = msgspec.field(default_factory=list)
    source_sheet: str | None = None
    source_range: str | None = None
    refresh_on_load: bool | None = None
    warnings: list[str] = msgspec.field(default_factory=list)


class PivotTablesSummary(msgspec.Struct, kw_only=True):
    """Pivot table metadata for a worksheet."""

    workbook: str
    sheet: str
    format: str
    pivot_tables: list[PivotTableInfo] = msgspec.field(default_factory=list)
    warnings: list[str] = msgspec.field(default_factory=list)


class PivotTablesMultiSheetResult(msgspec.Struct, kw_only=True):
    """Pivot table metadata for multiple worksheets."""

    workbook: str
    format: str
    sheets: list[PivotTablesSummary] = msgspec.field(default_factory=list)
    warnings: list[str] = msgspec.field(default_factory=list)


def get_pivot_tables_info(filepath: str, sheet_names: list[str]) -> str:
    """Read pivot table metadata for multiple sheets and return it as a JSON string."""

    path = Path(filepath)
    if not path.exists():
        raise ValidationError(f"File not found: {filepath}")

    if not sheet_names:
        raise ValidationError("At least one sheet name must be provided")

    suffix = path.suffix.lower()
    if suffix in XLSX_EXTENSIONS:
        result = _extract_xlsx_pivots_multi(path, sheet_names)
    elif suffix == ".xls":
        result = _extract_xls_pivots_multi(path, sheet_names)
    else:
        raise ValidationError(
            f"Unsupported workbook format '{path.suffix}'. "
            "Only .xlsx and .xls files are supported."
        )

    return msgjson.encode(result).decode("utf-8")


def _extract_xlsx_pivots_multi(path: Path, sheet_names: list[str]) -> PivotTablesMultiSheetResult:
    try:
        wb = load_workbook(path, data_only=True, keep_links=True, read_only=False)
    except InvalidFileException as exc:
        raise PivotError(f"Unable to open workbook: {exc}") from exc
    except Exception as exc:  # pragma: no cover - defensive
        raise PivotError(f"Failed to load workbook: {exc}") from exc

    try:
        summaries: list[PivotTablesSummary] = []
        global_warnings: list[str] = []
        missing_sheets: list[str] = []

        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                missing_sheets.append(sheet_name)
                continue

            try:
                ws = wb[sheet_name]
                pivot_defs = list(getattr(ws, "_pivots", []) or [])

                pivots: list[PivotTableInfo] = []
                warnings: list[str] = []

                for pivot in pivot_defs:
                    try:
                        pivots.append(_build_pivot_info(pivot))
                    except Exception as exc:  # pragma: no cover - narrow failure
                        name = getattr(pivot, "name", "unknown")
                        logger.exception("Failed to read pivot metadata for %s", name)
                        warnings.append(f"Failed to read pivot '{name}': {exc}")

                summaries.append(
                    PivotTablesSummary(
                        workbook=path.name,
                        sheet=sheet_name,
                        format="xlsx",
                        pivot_tables=pivots,
                        warnings=warnings,
                    )
                )
            except Exception as exc:
                logger.exception("Failed to process sheet '%s'", sheet_name)
                global_warnings.append(f"Failed to process sheet '{sheet_name}': {exc}")

        if missing_sheets:
            global_warnings.append(f"Sheet(s) not found: {', '.join(missing_sheets)}")

        return PivotTablesMultiSheetResult(
            workbook=path.name,
            format="xlsx",
            sheets=summaries,
            warnings=global_warnings,
        )
    finally:
        wb.close()


def _extract_xlsx_pivots(path: Path, sheet_name: str) -> PivotTablesSummary:
    """Legacy function for single sheet extraction."""
    result = _extract_xlsx_pivots_multi(path, [sheet_name])
    if not result.sheets:
        raise SheetError(f"Sheet '{sheet_name}' not found")
    return result.sheets[0]


def _build_pivot_info(pivot) -> PivotTableInfo:
    cache = getattr(pivot, "cache", None)
    if cache is None:
        raise PivotError("Pivot cache is missing")

    cache_fields = list(getattr(cache, "cacheFields", []) or [])

    def resolve_name(index: int | None) -> str | None:
        if index is None:
            return None
        try:
            field = cache_fields[index]
        except IndexError:
            return f"Field{index}"

        if field.name:
            return field.name

        shared = getattr(field, "sharedItems", None)
        if shared:
            for item in _iter_shared_items(shared):
                if item is not None:
                    return item
        return f"Field{index}"

    rows = [name for name in (resolve_name(getattr(f, "x", None)) for f in pivot.rowFields) if name]
    columns = [
        name for name in (resolve_name(getattr(f, "x", None)) for f in pivot.colFields) if name
    ]

    filters = [
        name
        for name in (
            resolve_name(getattr(page_field, "fld", None)) for page_field in pivot.pageFields
        )
        if name
    ]

    values: list[PivotValueInfo] = []
    for data_field in pivot.dataFields:
        source = resolve_name(getattr(data_field, "fld", None))
        values.append(
            PivotValueInfo(
                display_name=getattr(data_field, "name", None) or source,
                source_field=source,
                aggregation=_normalise_aggregation(getattr(data_field, "subtotal", None)),
            )
        )

    cache_source = getattr(cache, "cacheSource", None)
    worksheet_source = getattr(cache_source, "worksheetSource", None)
    source_sheet = getattr(worksheet_source, "sheet", None)
    source_range = getattr(worksheet_source, "ref", None)

    location = getattr(getattr(pivot, "location", None), "ref", None)
    refresh_on_load = getattr(cache, "refreshOnLoad", None)

    return PivotTableInfo(
        name=getattr(pivot, "name", None),
        location=location,
        rows=rows,
        columns=columns,
        values=values,
        filters=filters,
        source_sheet=source_sheet,
        source_range=source_range,
        refresh_on_load=refresh_on_load,
    )


def _extract_xls_pivots_multi(path: Path, sheet_names: list[str]) -> PivotTablesMultiSheetResult:
    global_warnings = [
        "Pivot metadata extraction for .xls files is limited. "
        "Convert the workbook to .xlsx for full details."
    ]

    summaries: list[PivotTablesSummary] = []
    missing_sheets: list[str] = []

    try:
        import xlrd  # Local import to avoid mandatory dependency when unused

        book = xlrd.open_workbook(path)
        available_sheets = book.sheet_names()

        for sheet_name in sheet_names:
            if sheet_name not in available_sheets:
                missing_sheets.append(sheet_name)
                continue

            pivot_names, scan_warnings = _scan_xls_pivot_names(path)
            global_warnings.extend(scan_warnings)

            pivot_infos: list[PivotTableInfo] = []
            for name in pivot_names:
                pivot_infos.append(
                    PivotTableInfo(
                        name=name,
                        warnings=["Only the pivot table name could be determined for this .xls file."],
                    )
                )

            summaries.append(
                PivotTablesSummary(
                    workbook=path.name,
                    sheet=sheet_name,
                    format="xls",
                    pivot_tables=pivot_infos,
                    warnings=[],
                )
            )
    except Exception as exc:
        raise PivotError(f"Failed to read .xls workbook: {exc}") from exc

    if missing_sheets:
        global_warnings.append(f"Sheet(s) not found: {', '.join(missing_sheets)}")

    return PivotTablesMultiSheetResult(
        workbook=path.name,
        format="xls",
        sheets=summaries,
        warnings=global_warnings,
    )


def _extract_xls_pivots(path: Path, sheet_name: str) -> PivotTablesSummary:
    """Legacy function for single sheet extraction."""
    result = _extract_xls_pivots_multi(path, [sheet_name])
    if not result.sheets:
        raise SheetError(f"Sheet '{sheet_name}' not found")
    return result.sheets[0]


def _scan_xls_pivot_names(path: Path) -> tuple[list[str], list[str]]:
    try:
        from olefile import OleFileIO
    except Exception as exc:  # pragma: no cover - import failure
        return [], [f"olefile not available: {exc}"]

    try:
        with OleFileIO(path) as ole:
            if not ole.exists("Workbook"):
                return [], ["Workbook stream not found in .xls file"]
            stream = ole.openstream("Workbook")
            data = stream.read()
    except Exception as exc:  # pragma: no cover - corrupted file
        return [], [f"Failed to read Workbook stream: {exc}"]

    marker = b"PivotTable"
    seen: set[str] = set()
    names: list[str] = []
    offset = 0
    length = len(data)

    while offset < length:
        idx = data.find(marker, offset)
        if idx == -1:
            break

        end = idx
        while end < length and data[end] not in (0x00,):
            end += 1

        raw = data[idx:end]
        try:
            name = raw.decode("latin1").strip()
        except UnicodeDecodeError:  # pragma: no cover - unlikely
            name = raw.decode("ascii", "ignore").strip()

        if not name:
            offset = end + 1
            continue

        match = re.match(r"(PivotTable\d+)", name)
        cleaned = match.group(1) if match else "".join(ch for ch in name if ch.isprintable())
        cleaned = cleaned.strip()

        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            names.append(cleaned)

        offset = end + 1

    if not names:
        return [], ["No pivot table identifiers were discovered in the .xls workbook stream."]

    return names, []


def _iter_shared_items(shared) -> Iterable[str | None]:
    """Return any textual items from a SharedItems definition."""

    for attr in ("s", "n", "m", "d", "longText"):
        sequence = getattr(shared, attr, None)
        if sequence is None:
            continue
        try:
            items = list(sequence)
        except TypeError:  # pragma: no cover - defensive
            continue
        for item in items:
            value = getattr(item, "v", None)
            if value is None and hasattr(item, "value"):
                value = getattr(item, "value", None)
            if value is not None:
                yield str(value)


def _normalise_aggregation(value: str | None) -> str | None:
    if not value:
        return None
    normalised = value.strip().lower()
    mapping = {
        "average": "average",
        "count": "count",
        "counta": "count",
        "max": "max",
        "min": "min",
        "product": "product",
        "stddev": "stddev",
        "stddevp": "stddevp",
        "sum": "sum",
        "var": "var",
        "varp": "varp",
    }
    return mapping.get(normalised, normalised)


